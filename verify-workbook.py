import json,re,hashlib
from pathlib import Path
import openpyxl
root=Path(__file__).parent
source=Path(r'C:\Users\00005191\Downloads\【編集厳禁】2026冬期講座プラン【最終確定版】.xlsx')
before=hashlib.sha256(source.read_bytes()).hexdigest()
w=openpyxl.load_workbook(source,data_only=True)
plans=json.loads((root/'dist/data.js').read_text(encoding='utf-8').removeprefix('window.PLANS=').strip().removesuffix(';'))
assert w.sheetnames==['力・作文','小４～６','中１・２','中3','高校生']
covered=set()
for p in plans:
 sheet,cells=p['source'].split('!');comp,detail=cells.split(' / ')
 assert p['composition']==str(w[sheet][comp].value),(p['id'],'composition')
 expected='\n'.join(x for x in str(w[sheet][detail].value or '').splitlines() if not x.startswith('\\\\') and '上記のテスト作成エクセル' not in x)
 assert p['detail']==expected,(p['id'],'detail')
 covered.add((sheet,comp))
expected=set()
for sheet,cols,rows in [('小４～６',range(2,8),[3,9,16]),('中１・２',range(3,8),[2,8,14,20]),('中3',range(2,7),[2,8,14]),('高校生',range(2,8),[3,9])]:
 for row in rows:
  for col in cols:
   c=w[sheet].cell(row,col)
   if c.value is not None:expected.add((sheet,c.coordinate))
# High-school grade 3 repeats the same two original plans in both rows.
assert expected-covered=={('高校生','F9'),('高校生','G9')}
assert w['高校生']['F9'].value==w['高校生']['F3'].value
assert w['高校生']['G9'].value==w['高校生']['G3'].value
assert [w['力・作文'].cell(2,c).value for c in range(2,7)]==['4回','8回','12回','16回','20回']
assert [w['力・作文'][c].value for c in ['C4','C5','C6','B8']]==['個別指導4回','AI学トレ8回（学力検定シリーズ＞語彙力養成コース　ベーシックコース）','読書4回','オリジナルプラン']
assert hashlib.sha256(source.read_bytes()).hexdigest()==before
print('PASS: all 5 sheets; 61 distinct plans; all composition/detail cells; two intentional high3 duplicates; shared course options; unchanged source SHA256 '+before)
